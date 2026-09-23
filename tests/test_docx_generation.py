"""docx 文书生成 — 模板/工具/图内流/交付端点 全链测试。"""
import re
import shutil
import zipfile
from pathlib import Path

TPL_DIR = Path(__file__).resolve().parents[1] / "data" / "doc_templates" / "complaint"


def _document_xml() -> str:
    """读取 template.docx 的 document.xml 原文。"""
    with zipfile.ZipFile(TPL_DIR / "template.docx") as z:
        return z.read("word/document.xml").decode("utf-8")


def test_template_docx_exists():
    """模板与源文件齐备: source.docx(原始附件)/template.docx(标签版)/fields.yaml。"""
    assert (TPL_DIR / "source.docx").exists()
    assert (TPL_DIR / "template.docx").exists()
    assert (TPL_DIR / "fields.yaml").exists()


def test_template_contains_all_text_tags():
    """fields.yaml 里全部 text/date 字段(含 _merge_into 附带字段)的 {{ f.* }} 必须落在模板里。"""
    import yaml

    spec = yaml.safe_load((TPL_DIR / "fields.yaml").read_text(encoding="utf-8"))
    xml = _document_xml()
    for f in spec["fields"]:
        if f["type"] == "choice":
            continue
        assert "{{ f." + f["key"] + " }}" in xml, f"缺文本标签 {f['key']}"


def test_template_contains_checkbox_tags():
    """抽查代表性勾选键(YAML 手写集的子集)。"""
    xml = _document_xml()
    for tag in ("c.plaintiff_gender_male", "c.defendant_gender_female",
                "c.property_has_none", "c.property_has_has",
                "c.debt_has_has", "c.agent_has_yes", "c.agent_has_no",
                "c.e_service_no", "c.e_service_method_other",
                "c.agent_scope_general",
                "c.compensation_damage", "c.visit_subject_plaintiff"):
        assert "{{ " + tag + " }}" in xml, f"缺勾选标签 {tag}"


def test_template_contains_all_handwritten_checkbox_tags():
    """fields.yaml 手写的每个 {{ c.* }}(以 replacement 实际所写为准)都必须落在模板里。"""
    import yaml

    spec = yaml.safe_load((TPL_DIR / "fields.yaml").read_text(encoding="utf-8"))
    xml = _document_xml()
    for f in spec["fields"]:
        for ckey in re.findall(r"\{\{\s*c\.([A-Za-z0-9_]+)\s*\}\}", f.get("replacement", "")):
            assert "{{ c." + ckey + " }}" in xml, f"缺勾选标签 {ckey} (字段 {f['key']})"
        if f.get("type") == "choice":
            assert re.search(r"\{\{ c\." + re.escape(f["key"]) + r"_[A-Za-z0-9_]+ \}\}", xml), \
                f"choice 字段 {f['key']} 无任何勾选标签"


# ---- Task 2: doc_templates 加载/对账 ----


def test_load_fields_complaint():
    from lawApp_LangGraph.doc_templates import load_fields
    fields = load_fields("complaint")
    assert fields, "complaint 模板应可加载"
    keys = {f["key"] for f in fields}
    assert {"plaintiff_name", "defendant_name", "fact_divorce_reason"} <= keys


def test_load_fields_missing_returns_empty():
    from lawApp_LangGraph.doc_templates import load_fields
    assert load_fields("defense") == []  # 本期无答辩状模板 → 空表(不抛)


def test_validate_template_complaint_ok():
    from lawApp_LangGraph.doc_templates import validate_template
    ok, missing = validate_template("complaint")
    assert ok, f"对账不应有缺口: {missing}"


def test_validate_template_detects_missing_tag(tmp_path):
    import lawApp_LangGraph.doc_templates as dt
    # 破坏副本: 复制 complaint 目录到 tmp, 删掉一个文本标签, 指向副本校验
    src = dt._TEMPLATE_ROOT / "complaint"
    dst = tmp_path / "broken"
    shutil.copytree(src, dst)
    doc = dst / "template.docx"
    # 先整体读入(避免同文件边读边写被 "w" 截断), 再重写破坏版
    with zipfile.ZipFile(doc) as zin:
        xml = zin.read("word/document.xml").decode("utf-8")
        payload = [(i.filename, zin.read(i.filename)) for i in zin.infolist()]
    xml = xml.replace("{{ f.plaintiff_name }}", "", 1)
    with zipfile.ZipFile(doc, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in payload:
            if name == "word/document.xml":
                data = xml.encode("utf-8")
            zout.writestr(name, data)
    ok, missing = dt._validate_dir(dst)
    assert not ok and "f.plaintiff_name" in missing


# ---- Task 3: generate_docx 工具 ----

import json as _json


def _run_tool(fields: dict, doc_type: str = "complaint", filename: str = "") -> dict:
    import asyncio
    from lawApp_LangGraph.tools.tools import generate_docx
    return asyncio.run(generate_docx.ainvoke(
        {"fields_json": _json.dumps(fields, ensure_ascii=False), "doc_type": doc_type, "filename": filename}
    ))


def _doc_text(docx_path) -> str:
    """取渲染产物全部可见文本: 表格嵌套逐层展开(模板是法院表格, 顶层段落无字段)。"""
    import docx  # docxtpl 带的 python-docx
    d = docx.Document(str(docx_path))
    parts = [p.text for p in d.paragraphs]

    def walk(table) -> None:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
                for nested in cell.tables:
                    walk(nested)

    for t in d.tables:
        walk(t)
    return "\n".join(parts)


def test_generate_docx_full_fill(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    fields = {"plaintiff_name": "张三", "plaintiff_gender": "男",
              "defendant_name": "李四", "defendant_gender": "女",
              "fact_divorce_reason": "感情破裂分居两年", "evidence_list": "结婚证、分居证明"}
    r = _run_tool(fields, filename="起诉状_AT-1.docx")
    assert r["status"] == "success" and r["docx_path"]
    import os
    assert os.path.exists(r["docx_path"])
    assert r["filled"] >= 6


def test_generate_docx_missing_text_becomes_pending(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    r = _run_tool({"plaintiff_name": "张三"})
    assert r["status"] == "success" and r["pending"] > 0
    # 产物里文本空位 = "待补充"
    assert "待补充" in _doc_text(r["docx_path"])


def test_generate_docx_choice_checkboxes(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    r = _run_tool({"plaintiff_gender": "男", "property_has": "有财产"})
    assert r["status"] == "success"
    text = _doc_text(r["docx_path"])
    assert "☑男 ☐女" in text
    assert "☐无财产" in text and "☑有财产" in text


def test_generate_docx_filename_sanitized(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    r = _run_tool({"plaintiff_name": "x"}, filename="../../evil/z..docx")
    assert r["status"] == "success"
    import os
    assert "evil" not in r["docx_path"].replace("\\", "/")
    assert os.path.dirname(os.path.abspath(r["docx_path"])) == str(tmp_path.resolve())


def test_generate_docx_merge_into_fields_rendered(tmp_path, monkeypatch):
    """I-1(R1): _merge_into 附带字段值须渲染(标签在宿主 replacement 内),
    缺失给"待补充"(spec D4); 仅 filled/pending 计数跳过。"""
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    # 传附带字段值: 渲染产物须含各值
    r = _run_tool({"plaintiff_work": "某公司", "plaintiff_duty": "经理",
                   "plaintiff_phone": "13800000000", "preservation_court": "海淀法院"})
    assert r["status"] == "success"
    text = _doc_text(r["docx_path"])
    for needle in ("经理", "13800000000", "海淀法院"):
        assert needle in text, f"附带字段值未被渲染: {needle}"
    # 只传宿主: 附带字段空位 = "待补充"(D4)
    r2 = _run_tool({"plaintiff_work": "某公司"})
    assert r2["status"] == "success"
    assert "职务：待补充" in _doc_text(r2["docx_path"])


def test_generate_docx_unknown_doctype_error(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    r = _run_tool({"a": "b"}, doc_type="defense")
    assert r["status"] == "error" and r["docx_path"] is None


def test_build_docx_template_reproducible(tmp_path):
    """I-2(deferred): 重建路径可复现 — 同源同 YAML 重跑 build 脚本,
    产物 document.xml 与已提交 template.docx 逐字节一致。"""
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    out = tmp_path / "template.docx"
    r = subprocess.run(
        [sys.executable, str(root / "scripts" / "build_docx_template.py"),
         "--src", str(TPL_DIR / "source.docx"),
         "--yaml", str(TPL_DIR / "fields.yaml"),
         "--out", str(out)],
        capture_output=True, text=True, timeout=120,
    )
    assert r.returncode == 0, r.stderr
    with zipfile.ZipFile(out) as z:
        rebuilt = z.read("word/document.xml")
    assert rebuilt == _document_xml().encode("utf-8")


# ---- Task 4: 图内流(docx_confirm interrupt 载荷 / 跳过 resume / 确认落盘) ----
# 驱动模式参照 tests/test_smoke.py 的 interrupt→Command(resume) 夹具:
# MemorySaver checkpointer + LLM 全替身(_stream_plan/_structured/finalize astream),
# _extract_doc_fields 打 AsyncMock; generate_docx 走真实 ToolNode 渲染。

import asyncio  # noqa: E402
import os  # noqa: E402
import sys  # noqa: E402
import types  # noqa: E402

if sys.platform == "win32":
    # psycopg async 在 Windows 需 SelectorEventLoop(对齐 test_dialogue_events)
    asyncio.set_event_loop_policy(asyncio.WindowsSelectorEventLoopPolicy())

_DOCX_T_PREFIX = "T-docx-test"

# _extract_doc_fields 打桩字段: 5 项已填 + 1 项空(待补充), 覆盖文本/choice
_DOCX_STUB_FIELDS = {
    "plaintiff_name": "张三",
    "defendant_name": "李四",
    "plaintiff_gender": "男",
    "defendant_gender": "女",
    "fact_divorce_reason": "感情破裂分居两年",
    "plaintiff_birth_date": "",  # 空 → field_preview pending/待补充
}


def _pg_ok() -> bool:
    """独立短连接探测 PG(不碰全局池, 对齐 test_dialogue_events)。"""

    async def _probe():
        from psycopg import AsyncConnection

        from lawApp_LangGraph.db import build_dsn

        conn = await AsyncConnection.connect(build_dsn(), autocommit=True)
        try:
            await conn.execute("SELECT 1")
        finally:
            await conn.close()

    try:
        asyncio.run(_probe())
        return True
    except Exception:
        return False


def _skip_if_no_pg():
    if not _pg_ok():
        import pytest

        pytest.skip("PG 不可用, 显式跳过(不 mock)")


def _docx_cleanup() -> None:
    """清理测试会话事件行(独立短连接)。"""

    async def _run():
        from psycopg import AsyncConnection

        from lawApp_LangGraph.db import build_dsn

        conn = await AsyncConnection.connect(build_dsn(), autocommit=True)
        try:
            await conn.execute(
                "DELETE FROM session_dialogue_events WHERE session_id LIKE %s",
                (_DOCX_T_PREFIX + "%",),
            )
        finally:
            await conn.close()

    try:
        asyncio.run(_run())
    except Exception:
        pass  # PG 不可用 → 后续用例各自 SKIP


def _count_events(sid: str, event_type: str) -> int:
    """查 session_dialogue_events 指定会话 + 类型行数。"""

    async def _run():
        from psycopg import AsyncConnection

        from lawApp_LangGraph.db import build_dsn

        conn = await AsyncConnection.connect(build_dsn(), autocommit=True)
        try:
            cur = await conn.execute(
                "SELECT count(*) FROM session_dialogue_events "
                "WHERE session_id = %s AND event_type = %s",
                (sid, event_type),
            )
            return (await cur.fetchone())[0]
        finally:
            await conn.close()

    return asyncio.run(_run())


def _build_docx_graph(monkeypatch):
    """装配图内流测试替身并 build_graph(MemorySaver)。

    LLM 全替身: risk/assess/replan_check 放行、finalize 出固定文本;
    planner 经 _stream_plan 返回单步 generate_docx 计划; 抽取打 AsyncMock。
    """
    from unittest.mock import AsyncMock

    import lawApp_LangGraph.LangGraph_lawApp as app
    from langchain_core.runnables import Runnable
    from langgraph.checkpoint.memory import MemorySaver

    verdict = types.SimpleNamespace(
        high_risk=False, question_category="marriage_legal", applicable=True,
        element_updates=[], na_keys=[], promote_keys=[], questions=[], done=True,
        needs_replan=False, reason="", insufficient_reason="none",
    )

    class _Chain(Runnable):
        """`PromptTemplate | _structured(...)` 决策链替身。"""

        def invoke(self, _inp, config=None, **_kw):
            return verdict

        async def ainvoke(self, _inp, config=None, **_kw):
            return verdict

    class _StubLLM(Runnable):
        """executor/finalize LLM 替身: 结构化链回固定 verdict, astream 出终答文本。"""

        def with_structured_output(self, _schema, **_kw):
            return _Chain()

        def bind_tools(self, _tools):
            return self

        def invoke(self, _msgs, config=None, **_kw):
            return verdict

        async def ainvoke(self, _msgs, config=None, **_kw):
            return verdict

        async def astream(self, _prompt, config=None, **_kw):
            yield types.SimpleNamespace(content="文书终答测试")

    async def _fake_stream_plan(_prompt_text, _source, _config):
        plan = app.PlanSchema.model_validate(
            {
                "reasoning": ["文书起草"],
                "plan": [
                    {
                        "step_id": 1,
                        "description": "按起诉状模板生成 Word 文书",
                        "tool_name": "generate_docx",
                    }
                ],
            }
        )
        return plan, []

    monkeypatch.setattr(app, "get_executor_llm", lambda: _StubLLM())
    monkeypatch.setattr(app, "get_planner_llm", lambda: _StubLLM())
    monkeypatch.setattr(app, "_stream_plan", _fake_stream_plan)
    monkeypatch.setattr(
        app, "_extract_doc_fields", AsyncMock(return_value=dict(_DOCX_STUB_FIELDS))
    )
    return app.build_graph(checkpointer=MemorySaver())


def _docx_invoke_input() -> dict:
    return {
        "query": "我要起诉离婚, 请帮我生成起诉状",
        "mode": "assistant",
        "doc_type": "complaint",
    }


def test_docx_confirm_interrupt_payload(monkeypatch):
    """executor docx 步骤: 触发 docx_confirm, 载荷含 field_preview 与二选选项。"""
    g = _build_docx_graph(monkeypatch)
    cfg = {"configurable": {"thread_id": _DOCX_T_PREFIX + "-payload"},
           "recursion_limit": 40}

    async def run():
        result = await g.ainvoke(_docx_invoke_input(), config=cfg)
        assert not result.get("final_answer"), "interrupt 前不得出终答"

        snap = await g.aget_state(cfg)
        assert snap.next, "应停在 docx_confirm interrupt"
        intr = next(iter(snap.interrupts), None)
        assert intr is not None, "pending interrupts 应含 docx_confirm"
        p = intr.value
        assert p["type"] == "docx_confirm"
        assert "确认生成吗" in p["message"]
        assert "已填 5 项" in p["message"] and "待补充 1 项" in p["message"]
        assert p["options"] == [
            {"value": "确认", "label": "确认生成 Word 文书"},
            {"value": "跳过", "label": "跳过该步骤"},
        ]
        prev = {item["key"]: item for item in p["field_preview"]}
        assert prev["plaintiff_name"] == {
            "key": "plaintiff_name", "label": "原告姓名", "value": "张三",
            "critical": True, "status": "filled",
        }
        assert prev["plaintiff_birth_date"]["status"] == "pending"
        assert prev["plaintiff_birth_date"]["value"] == "待补充"

    asyncio.run(run())


def test_docx_confirm_skip_resume_no_docx(monkeypatch, tmp_path):
    """resume 跳过: finalize 正常收尾, state 无 docx_path, 无 docx_generated 事件。"""
    _skip_if_no_pg()
    _docx_cleanup()
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    g = _build_docx_graph(monkeypatch)
    sid = _DOCX_T_PREFIX + "-skip"
    cfg = {"configurable": {"thread_id": sid}, "recursion_limit": 40}

    async def run():
        from langgraph.types import Command

        await g.ainvoke(_docx_invoke_input(), config=cfg)  # 停在 docx_confirm
        result = await g.ainvoke(Command(resume=False), config=cfg)
        assert result.get("final_answer") == "文书终答测试"
        assert result.get("docx_path") is None, "跳过不得生成 docx"
        assert result.get("docx_confirmed") is True, "跳过后确认位落定(防再问)"
        assert not list(tmp_path.glob("*.docx")), "跳过不得落盘任何 docx"

    try:
        asyncio.run(run())
        # 落库: docx_confirm(跳过)一行, docx_generated 零行
        assert _count_events(sid, "docx_confirm") == 1
        assert _count_events(sid, "docx_generated") == 0
    finally:
        _docx_cleanup()


def test_docx_confirm_yes_resume_generates_file(monkeypatch, tmp_path):
    """resume 确认: generate_docx 真实渲染(docxtpl), docx_path 进 state,
    docx_confirm + docx_generated 两事件落库(PG 查 session_dialogue_events)。"""
    _skip_if_no_pg()
    _docx_cleanup()
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    g = _build_docx_graph(monkeypatch)
    sid = _DOCX_T_PREFIX + "-yes"
    cfg = {"configurable": {"thread_id": sid}, "recursion_limit": 40}

    async def run():
        from langgraph.types import Command

        await g.ainvoke(_docx_invoke_input(), config=cfg)  # 停在 docx_confirm
        result = await g.ainvoke(Command(resume=True), config=cfg)
        assert result.get("final_answer") == "文书终答测试"
        assert result.get("docx_path"), "确认后 docx_path 应进 state"
        assert os.path.exists(result["docx_path"]), "docx 必须真实落盘"
        assert result.get("docx_confirmed") is True
        # 渲染内容抽查: 抽取字段真实进了 Word
        text = _doc_text(result["docx_path"])
        assert "张三" in text and "李四" in text
        assert "感情破裂分居两年" in text
        return result

    try:
        result = asyncio.run(run())
        # 落库: docx_confirm(确认) + docx_generated 各至少一行
        assert _count_events(sid, "docx_confirm") == 1
        assert _count_events(sid, "docx_generated") == 1
        assert os.path.basename(result["docx_path"]).startswith("起诉状_")
    finally:
        _docx_cleanup()


# ---- R1: 抽取失败 retry 一次(M-3) + 直调空字段守卫(I-2) ----


def test_docx_extract_double_failure_marks_step_failed(monkeypatch, tmp_path):
    """M-3(spec §7): 抽取两次全失败 → 步骤 failed(retry_count+1), 无 docx
    落盘、无 docx_confirm/docx_generated 事件, 图正常收尾(终答仍出)。"""
    _skip_if_no_pg()
    _docx_cleanup()
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    from unittest.mock import AsyncMock

    import lawApp_LangGraph.LangGraph_lawApp as app

    g = _build_docx_graph(monkeypatch)
    # 两次抽取全抛(AsyncMock side_effect 顺序消费)
    monkeypatch.setattr(
        app,
        "_extract_doc_fields",
        AsyncMock(side_effect=[RuntimeError("抽取失败1"), RuntimeError("抽取失败2")]),
    )
    sid = _DOCX_T_PREFIX + "-fail2"
    cfg = {"configurable": {"thread_id": sid}, "recursion_limit": 40}

    async def run():
        # 抽取在 interrupt 之前失败 → 单次 ainvoke 直接跑完, 无需 resume
        result = await g.ainvoke(_docx_invoke_input(), config=cfg)
        assert result.get("final_answer") == "文书终答测试", "图应正常收尾"
        assert result["plan"][0].status == "failed", "两次失败后步骤必须标 failed"
        assert result["plan"][0].retry_count == 1
        assert result.get("error"), "应写入步骤 error"
        assert result.get("docx_confirmed") is True, "失败路径确认位落定(防再问)"
        assert result.get("docx_path") is None

    try:
        asyncio.run(run())
        assert not list(tmp_path.glob("*.docx")), "失败路径不得落盘 docx"
        assert _count_events(sid, "docx_generated") == 0
        assert _count_events(sid, "docx_confirm") == 0, "未到 interrupt, 无确认事件"
    finally:
        _docx_cleanup()


def test_docx_direct_call_empty_fields_guarded(monkeypatch):
    """I-2: 抽取失败遗留状态(docx_confirmed=True + doc_fields={})再进直调
    分支(replanner 重出 generate_docx 步) → 守卫生效: 不产 tool_calls,
    步骤按跳过标 done 推进, 不渲染空白文书(违 D3)。"""
    import lawApp_LangGraph.LangGraph_lawApp as app
    from lawApp_LangGraph.state import AgentState, PlanStep

    state = AgentState(
        query="我要起诉离婚, 请帮我生成起诉状",
        mode="assistant",
        doc_type="complaint",
        plan=[
            PlanStep(
                step_id=1,
                description="按起诉状模板生成 Word 文书",
                tool_name="generate_docx",
            )
        ],
        current_step_index=0,
        # 抽取失败路径的遗留状态: 确认位已置但无字段
        doc_fields={},
        docx_confirmed=True,
    )
    out = asyncio.run(app.executor_node(state, None))

    assert "messages" not in out, "不得产 AIMessage tool_calls(不调工具)"
    assert out["plan"][0].status == "done", "步骤按跳过标 done"
    assert out["current_step_index"] == 1, "步骤索引推进"
    assert out["docx_confirmed"] is True


# ---- Task 5: API 层(normalize 集合 / 交付端点 / dialogue docx 键) ----
# 端点测试走 TestClient(仓库惯例, 路由无 /api 前缀), 造数用 dialogue_log.log_event
# 直插 session_dialogue_events(对齐 test_dialogue_events 的既有造数模式)。

_DOCX_MIME = (
    "application/vnd.openxmlformats-officedocument"
    ".wordprocessingml.document"
)


def test_docx_confirm_normalize_short_words():
    """docx_confirm 进归一化集合: 短词直接命中(不走 LLM, 无网络)。"""
    from lawApp_LangGraph.FastAPI.utils import normalize_resume

    assert asyncio.run(
        normalize_resume("docx_confirm", "确认", "生成文书?")
    ) is True
    assert asyncio.run(
        normalize_resume("docx_confirm", "跳过", "生成文书?")
    ) is False


def test_docx_latest_endpoint_404_when_none():
    """无 docx_generated 事件的会话 → 404(三态之一)。"""
    _skip_if_no_pg()
    from fastapi.testclient import TestClient

    from lawApp_LangGraph.FastAPI.api import app

    with TestClient(app) as client:
        # 格式合法但从不存在的会话 → 404 而非 500
        r = client.get("/sessions/AT-20990101-000000-999/docx/latest")
        assert r.status_code == 404
        assert "尚未生成" in r.json()["detail"]


def test_docx_latest_endpoint_serves_file(tmp_path, monkeypatch):
    """真实链路: 造 docx_generated 事件 + 落盘文件 → 200 + docx MIME + 文件流。"""
    _skip_if_no_pg()
    _docx_cleanup()
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    from lawApp_LangGraph import dialogue_log

    sid = _DOCX_T_PREFIX + "-api-latest"
    f = tmp_path / "起诉状_test.docx"
    f.write_bytes(b"PK\x03\x04fake-docx-payload")
    dialogue_log.log_event(
        sid, "docx_generated",
        {"docx_path": str(f), "filled": 6, "pending": 1},
    )
    try:
        from fastapi.testclient import TestClient

        from lawApp_LangGraph.FastAPI.api import app

        with TestClient(app) as client:
            r = client.get(f"/sessions/{sid}/docx/latest")
            assert r.status_code == 200, r.text
            assert r.headers["content-type"].startswith(_DOCX_MIME)
            assert r.content[:2] == b"PK"
            assert r.content == b"PK\x03\x04fake-docx-payload"
    finally:
        _docx_cleanup()


def test_docx_latest_rejects_path_escape(tmp_path, monkeypatch):
    """docx_generated 事件里 path 指向 DOCX_OUTPUT_DIR 外 → 404(白名单校验)。"""
    _skip_if_no_pg()
    _docx_cleanup()
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path / "out"))
    from lawApp_LangGraph import dialogue_log

    sid = _DOCX_T_PREFIX + "-api-escape"
    outside = tmp_path / "evil.docx"  # 在 DOCX_OUTPUT_DIR 之外
    outside.write_bytes(b"PK\x03\x04secret")
    dialogue_log.log_event(
        sid, "docx_generated",
        {"docx_path": str(outside), "filled": 1, "pending": 0},
    )
    try:
        from fastapi.testclient import TestClient

        from lawApp_LangGraph.FastAPI.api import app

        with TestClient(app) as client:
            r = client.get(f"/sessions/{sid}/docx/latest")
            assert r.status_code == 404, "路径越界必须被白名单拒绝"
            assert "路径非法" in r.json()["detail"] or "不存在" in r.json()["detail"]
    finally:
        _docx_cleanup()


def test_docx_latest_404_when_file_missing(tmp_path, monkeypatch):
    """事件指向的文件已被清理(磁盘丢失) → 404(三态之一)。"""
    _skip_if_no_pg()
    _docx_cleanup()
    monkeypatch.setenv("DOCX_OUTPUT_DIR", str(tmp_path))
    from lawApp_LangGraph import dialogue_log

    sid = _DOCX_T_PREFIX + "-api-missing"
    # 只造事件不落盘文件(模拟产物被运维清理)
    dialogue_log.log_event(
        sid, "docx_generated",
        {"docx_path": str(tmp_path / "起诉状_gone.docx"), "filled": 3, "pending": 0},
    )
    try:
        from fastapi.testclient import TestClient

        from lawApp_LangGraph.FastAPI.api import app

        with TestClient(app) as client:
            r = client.get(f"/sessions/{sid}/docx/latest")
            assert r.status_code == 404
    finally:
        _docx_cleanup()


def test_dialogue_aggregate_includes_docx_key(tmp_path):
    """dialogue aggregate: docx_generated → docx 键取最后一条; 无事件 → None。"""
    _skip_if_no_pg()
    _docx_cleanup()
    from lawApp_LangGraph import dialogue_log

    sid = _DOCX_T_PREFIX + "-agg-docx"
    f = tmp_path / "起诉状_agg.docx"
    f.write_bytes(b"PK")
    dialogue_log.log_event(
        sid, "docx_generated",
        {"docx_path": str(tmp_path / "old.docx"), "filled": 1, "pending": 9},
    )
    dialogue_log.log_event(
        sid, "docx_generated",
        {"docx_path": str(f), "filled": 6, "pending": 1},
    )

    async def _run():
        from lawApp_LangGraph.db import close_pool

        doc = await dialogue_log.aggregate_dialogue(sid)
        empty = await dialogue_log.aggregate_dialogue(_DOCX_T_PREFIX + "-agg-none")
        await close_pool()
        return doc, empty

    try:
        doc, empty = asyncio.run(_run())
        # 多条 docx_generated → 取最后一条(seq 最大)
        assert doc["docx"] == {
            "path": str(f), "filled": 6, "pending": 1,
        }
        # 无 docx 事件的会话 → docx=None(前端据此隐藏下载入口)
        assert empty["docx"] is None
    finally:
        _docx_cleanup()
